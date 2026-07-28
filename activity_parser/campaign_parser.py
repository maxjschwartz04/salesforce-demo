"""
Parses a Salesforce Campaign Member report export (.xlsx) into per-company
marketing engagement records — webinar RSVPs/attendance, website form
fills, etc. Distinct from sales activity: this is top-of-funnel marketing
engagement, not a rep's own outreach, and isn't matched into the revival
library or used for staleness math.

This export is GROUPED by week of "Member First Associated Date" — the
date only appears on the first row of each week's group, and is blank for
every row after it until the next group starts. That's Salesforce's report
display convention, not missing data — parse_campaign_report() forward-
fills the date group onto every row so each record carries its own value.
The result is only ever a week RANGE (e.g. "7/19/2026 - 7/25/2026"), not an
exact date — that's the finest granularity this particular export gives.

Usage:
    python campaign_parser.py campaigns.xlsx -o campaign_engagement.json
"""

import argparse
import json
import re
import sys
from datetime import datetime

import openpyxl

_DATE_GROUP_START_PATTERN = re.compile(r"(\d{1,2}/\d{1,2}/\d{4})")


def _parse_date_group_start(date_group):
    """date_group is a week-range string like '7/19/2026 - 7/25/2026' — this
    is the finest granularity the export gives us, so sorting/recency uses
    the range's start date, not an exact day."""
    if not date_group:
        return None
    match = _DATE_GROUP_START_PATTERN.search(date_group)
    if not match:
        return None
    try:
        return datetime.strptime(match.group(1), "%m/%d/%Y")
    except ValueError:
        return None


def _find_header(rows):
    for i, row in enumerate(rows):
        col = {}
        for j, v in enumerate(row):
            if v is None:
                continue
            key = str(v).strip()
            col[key] = j
        if "Company" in col and "Campaign Name" in col:
            return i, col
    raise ValueError("Could not find a header row containing 'Company' and 'Campaign Name'")


def _find_date_group_column(col):
    for key in col:
        if key.startswith("Member First Associated Date"):
            return col[key]
    return None


def parse_campaign_report(path):
    """Returns a flat list of {date_group, member_type, first_name,
    last_name, title, email, company, campaign_name, member_status,
    campaign_type} dicts, one per campaign membership row."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))

    header_idx, col = _find_header(rows)
    date_group_idx = _find_date_group_column(col)

    def get(row, name):
        idx = col.get(name)
        return row[idx] if idx is not None else None

    records = []
    current_date_group = None
    for row in rows[header_idx + 1 :]:
        company = get(row, "Company")
        if company is None:
            continue  # blank row, or a footer/subtotal line

        if date_group_idx is not None and row[date_group_idx]:
            current_date_group = row[date_group_idx]

        records.append(
            {
                "date_group": current_date_group,
                "member_type": get(row, "Member Type"),
                "first_name": get(row, "First Name"),
                "last_name": get(row, "Last Name"),
                "title": get(row, "Title"),
                "email": get(row, "Email"),
                "company": company,
                "campaign_name": get(row, "Campaign Name"),
                "member_status": get(row, "Member Status"),
                "campaign_type": get(row, "Campaign Type"),
            }
        )

    return records


def group_by_company(records):
    """Returns {company: [records]}."""
    grouped = {}
    for r in records:
        grouped.setdefault(r["company"], []).append(r)
    return grouped


def get_recent_engagements(company, grouped, top_n=3):
    """Most recent campaign engagements for a company, newest first. Exact
    company-name match only — deliberately NOT fuzzy-matched like
    opportunity_parser's find_account_status. "Company" on a Lead record is
    self-typed by whoever filled out a form, not a real Account name, so
    it's noisier; fuzzy-matching that against a canonical account name
    risks pairing the wrong company rather than just missing a real one."""
    engagements = grouped.get(company, [])
    dated = [(e, _parse_date_group_start(e["date_group"])) for e in engagements]
    dated = [(e, d) for e, d in dated if d is not None]
    dated.sort(key=lambda x: x[1], reverse=True)
    return [e for e, _ in dated[:top_n]]


def top_contacts(company, grouped, limit=3):
    """The most relevant real contacts on file for a company, so a rep
    knows who to reach out to without going back to Salesforce. The same
    person often shows up across several campaign touches, so this dedupes
    by email and ranks by that contact's OWN most recent engagement (same
    recency-first reasoning as get_recent_engagements) rather than by raw
    touch count, which would favor someone contacted often a long time ago
    over someone who just engaged.

    Exact company-name match only, same reasoning as get_recent_engagements.
    Contacts with no email on file are skipped -- nothing to link to.

    Returns [{first_name, last_name, title, email, engagement_count,
    most_recent_date_group}, ...], most-recently-engaged first."""
    engagements = grouped.get(company, [])

    by_email = {}
    for e in engagements:
        email = e.get("email")
        if not email:
            continue
        key = email.strip().lower()
        entry = by_email.setdefault(
            key,
            {
                "first_name": None,
                "last_name": None,
                "title": None,
                "email": email,
                "engagement_count": 0,
                "_most_recent_dt": None,
                "most_recent_date_group": None,
            },
        )
        entry["engagement_count"] += 1

        touch_dt = _parse_date_group_start(e.get("date_group"))
        if touch_dt is not None and (entry["_most_recent_dt"] is None or touch_dt > entry["_most_recent_dt"]):
            entry["_most_recent_dt"] = touch_dt
            entry["most_recent_date_group"] = e.get("date_group")
            # Keep whichever name/title came with the most recent touch --
            # a title on file from years ago may no longer be accurate.
            entry["first_name"] = e.get("first_name")
            entry["last_name"] = e.get("last_name")
            entry["title"] = e.get("title")

    contacts = list(by_email.values())
    contacts.sort(key=lambda c: c["_most_recent_dt"] or datetime.min, reverse=True)
    for c in contacts:
        del c["_most_recent_dt"]
    return contacts[:limit]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("report_path", help="Path to the Campaign Member report export (.xlsx)")
    parser.add_argument("-o", "--output", help="Write JSON to this file instead of stdout")
    args = parser.parse_args()

    records = parse_campaign_report(args.report_path)
    grouped = group_by_company(records)
    print(f"parsed {len(records)} campaign membership rows across {len(grouped)} companies", file=sys.stderr)

    output = json.dumps(grouped, indent=2, ensure_ascii=False)
    if args.output:
        with open(args.output, "w", encoding="utf-8") as f:
            f.write(output)
    else:
        print(output)


if __name__ == "__main__":
    sys.exit(main())
